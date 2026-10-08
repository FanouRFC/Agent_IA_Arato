"""Rapports PDF (reportlab) et Excel (openpyxl) – bonus, à la demande."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


def _rows(tickets, tracking):
    return [[f"T-{i:03d}", t["titre"], t["statut"], f"{t['avancement_declare']} %", tracking[i]["libelle"],
             t.get("motif_retard_courant") or t.get("motif_attente_courant") or "—"] for i, t in tickets.items()]


HEAD = ["Ticket", "Titre", "Statut CRM", "Avancement déclaré", "État de l'échéance", "Motif documenté"]


def pdf(synthese: str, tickets, tracking) -> bytes:
    buf, st = io.BytesIO(), getSampleStyleSheet()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=30, rightMargin=30)
    cell = st["BodyText"].clone("c", fontSize=7, leading=9)
    data = [HEAD] + [[Paragraph(str(c), cell) for c in r] for r in _rows(tickets, tracking)]
    tb = Table(data, repeatRows=1, colWidths=[40, 120, 55, 55, 110, 100])
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F4C47")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, 0), 7),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    doc.build([Paragraph("Rapport de suivi – Arato", st["Title"]), Paragraph("Synthèse de la situation", st["Heading2"]),
               Paragraph(synthese, st["BodyText"]), Spacer(1, 12), Paragraph("Tickets, retards et motifs", st["Heading2"]), tb])
    return buf.getvalue()


def xlsx(synthese: str, tickets, tracking) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Suivi"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in _rows(tickets, tracking):
        ws.append(r)
    for col, w in zip("ABCDEF", (10, 40, 14, 18, 36, 36)):
        ws.column_dimensions[col].width = w
    s = wb.create_sheet("Synthèse")
    s["A1"], s["A1"].font = "Synthèse de la situation", Font(bold=True)
    s["A2"] = synthese
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
